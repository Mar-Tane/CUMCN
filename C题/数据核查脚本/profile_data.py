# -*- coding: utf-8 -*-
import openpyxl, os
import statistics as st

BASE = os.path.join("E:\\", "数模", "26国赛", "26 N AI assist'", "C题", "附件")

def load(path, sheet=0):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = sheet if isinstance(sheet, str) else wb.sheetnames[sheet]
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows

# ---------- 附件1 ----------
r = load(os.path.join(BASE, "附件1.xlsx"))
print("附件1 header:", r[0], " nrows:", len(r))
print("附件1 first:", r[1])
print("附件1 last3:", r[-3:])
times = [x[0] for x in r[1:]]
price = [float(x[1]) for x in r[1:]]
load1 = [float(x[2]) for x in r[1:]]
pv1 = [float(x[3]) for x in r[1:]]
print("附件1 时段数:", len(times))
print("电价 min/max/mean: %.4f %.4f %.4f" % (min(price), max(price), sum(price)/len(price)))
print("负载 min/max/mean: %.2f %.2f %.2f" % (min(load1), max(load1), sum(load1)/len(load1)))
print("光伏 min/max/日总量kWh: %.2f %.2f %.2f" % (min(pv1), max(pv1), sum(pv1)/6.0))
print("负载日总量kWh: %.2f" % (sum(load1)/6.0))
print("光伏>0 时段数:", sum(1 for v in pv1 if v > 0))
argmax = max(range(len(price)), key=lambda i: price[i]); argmin = min(range(len(price)), key=lambda i: price[i])
print("电价最高时段:", times[argmax], price[argmax], " 最低时段:", times[argmin], price[argmin])

# ---------- 附件2 ----------
def load_wide(path, sheet):
    rows = load(path, sheet)
    hdr = rows[0]
    data = {}
    for row in rows[1:]:
        d = str(row[0])[:10]
        vals = []
        for v in row[1:]:
            vals.append(None if v is None else float(v))
        data[d] = vals
    return hdr, data

h2a, load2 = load_wide(os.path.join(BASE, "附件2.xlsx"), "小区负载")
h2b, pv2 = load_wide(os.path.join(BASE, "附件2.xlsx"), "光伏发电实际功率")
h4, price4 = load_wide(os.path.join(BASE, "附件4.xlsx"), "Sheet1")
print("\n附件2 负载 天数:", len(load2), " 每天点数:", len(next(iter(load2.values()))))
print("附件2 header 首/末时间列:", h2a[1], h2a[-1])
print("附件2 日期范围:", min(load2), "~", max(load2))
allL = [v for d in load2 for v in load2[d]]
allP = [v for d in pv2 for v in pv2[d]]
print("全年负载 min/max/mean: %.2f %.2f %.2f" % (min(allL), max(allL), sum(allL)/len(allL)))
print("全年光伏 min/max/mean: %.2f %.2f %.2f" % (min(allP), max(allP), sum(allP)/len(allP)))
print("光伏负值个数:", sum(1 for v in allP if v < 0), " 负载缺失:", sum(1 for v in allL if v is None))
print("负载峰值日:", max(load2, key=lambda d: max(load2[d])), max(max(v) for v in load2.values()))
dailypv = {d: sum(v)/6.0 for d, v in pv2.items()}
print("日光伏电量 min/max: %.1f %.1f kWh" % (min(dailypv.values()), max(dailypv.values())))
print("最大光伏功率出现日:", max(pv2, key=lambda d: max(pv2[d])), max(max(v) for v in pv2.values()))

print("\n附件4 电价 天数:", len(price4), " 点数/天:", len(next(iter(price4.values()))))
allpr = [v for d in price4 for v in price4[d]]
print("电价 min/max/mean: %.4f %.4f %.4f" % (min(allpr), max(allpr), sum(allpr)/len(allpr)))
neg = sum(1 for v in allpr if v < 0)
print("负电价个数:", neg, " 电价<0.1个数:", sum(1 for v in allpr if v < 0.1))
print("附件4 header 首/末:", h4[1], h4[-1])

# ---------- 附件3 ----------
r3 = load(os.path.join(BASE, "附件3.xlsx"))
print("\n附件3 header:", r3[0])
print("附件3 行数:", len(r3), " 列数:", len(r3[0]))
print("附件3 r1:", r3[1][:8])
print("附件3 r2:", r3[2][:8])
print("附件3 r3:", r3[3][:8])
print("附件3 r4:", r3[4][:8])
print("附件3 末行:", r3[-1][:8])
# 预报时刻集合
mom = [r3[i][1] for i in range(1, len(r3))]
print("预报时刻取值:", sorted(set(str(m) for m in mom)))
# 预报列数
print("预报列标题:", r3[0][2:])
# 检查 2025.3.20 等指定日
target = ["2025-3-20", "2025-6-21", "2025-9-23", "2025-12-21"]
for i in range(1, len(r3)):
    if str(r3[i][0]) in target:
        print("  找到", r3[i][0], r3[i][1], "首值", r3[i][2])
