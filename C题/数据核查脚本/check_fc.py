# -*- coding: utf-8 -*-
import openpyxl, os

BASE = os.path.join("E:\\", "数模", "26国赛", "26 N AI assist'", "C题", "附件")

def rows(path, sheet=0):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = sheet if isinstance(sheet, str) else wb.sheetnames[sheet]
    ws = wb[name]
    out = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out

# 附件2 光伏实际（10min）
r = rows(os.path.join(BASE, "附件2.xlsx"), "光伏发电实际功率")
hdr = [str(h) for h in r[0][1:]]
pv_act = {}
for row in r[1:]:
    pv_act[str(row[0])[:10]] = [float(v) for v in row[1:]]
print("附件2 时间列(前5):", hdr[:5], " 末列:", hdr[-1])

# 把10min聚合成整点均值(kW): 第k小时 = 索引 6k..6k+5  (00:10...01:00)
# 注意首列是 00:10，末列是 0:00+1 (次日0:00)
def hourly(v):
    # v: 144个10min点，对应 00:10,00:20,...,23:50,24:00
    # 整点小时 h (0..23) 用 v[6h : 6h+6] 表示 h:10 ~ h+1:00
    return [sum(v[6*h:6*h+6])/6.0 for h in range(24)]

# 附件3 预报
r3 = rows(os.path.join(BASE, "附件3.xlsx"))
fc = {}
cur = None
for row in r3[1:]:
    d = str(row[0]).strip()
    if d:
        cur = d
    mom = str(row[1])
    fc[(cur, mom)] = [float(x) for x in row[2:]]

# 对齐检验：2025-1-1 0:00 预报 vs 当天实际
import datetime
def norm(d):
    y, m, dd = [int(x) for x in d.split("-")]
    return datetime.date(y, m, dd)

for day in ["2025-1-1", "2025-3-20", "2025-6-21"]:
    key = None
    for k in pv_act:
        if norm(k) == norm(day):
            key = k
    act = hourly(pv_act[key])
    f0 = fc.get((day if day in dict(((a,b),1) for a,b in fc) else None, "0:00"))
    # 直接按原字符串
    f0 = None
    for (dd, mm), v in fc.items():
        if norm(dd) == norm(day) and mm == "0:00":
            f0 = v
    print("\n日期", day)
    print("  实际整点(前10):", [round(x, 1) for x in act[:10]])
    print("  预报0:00(前10):", [round(x, 1) for x in f0[:10]])
    # 两种对齐假设的RMSE
    def rmse(f, a, shift):
        # f[k] 对应 实际小时 k+shift
        s = 0; n = 0
        for k in range(24):
            j = k + shift
            if 0 <= j < 24:
                s += (f[k]-a[j])**2; n += 1
        return (s/n)**0.5
    print("  RMSE 假设A(预报k小时=实际第k小时, shift=0): %.1f" % rmse(f0, act, 0))
    print("  RMSE 假设B(预报k小时=实际第k+1小时, shift=1): %.1f" % rmse(f0, act, 1))

# 预报误差统计
diffs = []
for (d, m), v in fc.items():
    if m != "0:00":
        continue
    key = None
    for k in pv_act:
        if norm(k) == norm(d):
            key = k
    if key is None:
        continue
    act = hourly(pv_act[key])
    for k in range(24):
        diffs.append(v[k] - act[k])
n = len(diffs)
mean = sum(diffs)/n
sd = (sum((x-mean)**2 for x in diffs)/n)**0.5
print("\n0:00预报 - 实际 误差: n=%d 均值=%.2f kW 标准差=%.2f kW" % (n, mean, sd))
scale = sum(hourly(pv_act[k])[h] for k in pv_act for h in range(24))
print("误差标准差 / 全年小时光伏均值 = %.3f" % (sd / (sum(sum(hourly(v)) for v in pv_act.values())/(365*24))))

# 各预报时刻误差
for m in ["0:00", "6:00", "12:00", "18:00"]:
    dd = []
    for (d, mm), v in fc.items():
        if mm != m: continue
        key = None
        for k in pv_act:
            if norm(k) == norm(d): key = k
        if key is None: continue
        act = hourly(pv_act[key])
        # 预报k小时 -> 实际 (发布时刻+k) 小时, 可能跨日，简化只取当日
        off = {"0:00": 0, "6:00": 6, "12:00": 12, "18:00": 18}[m]
        for k in range(24):
            j = off + k
            if j < 24:
                dd.append(v[k] - act[j])
    if dd:
        mu = sum(dd)/len(dd); s = (sum((x-mu)**2 for x in dd)/len(dd))**0.5
        print("  预报时刻 %s: n=%d 偏差均值=%.2f 标准差=%.2f" % (m, len(dd), mu, s))
