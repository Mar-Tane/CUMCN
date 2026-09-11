# -*- coding: utf-8 -*-
import openpyxl, os, sys

BASE = os.path.join("E:\\", "数模", "26国赛", "26 N AI assist'", "C题", "附件")
files = ["附件1.xlsx", "附件2.xlsx", "附件3.xlsx", "附件4.xlsx"]
res = ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"]

def peek(path, maxr=4, maxc=16):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for s in wb.sheetnames:
        ws = wb[s]
        out.append("  [sheet] %s  rows=%s cols=%s" % (s, ws.max_row, ws.max_column))
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= maxr:
                break
            vals = ["" if v is None else str(v)[:20] for v in row[:maxc]]
            out.append("    r%-2d %s" % (i, vals))
    wb.close()
    return "\n".join(out)

for f in files:
    p = os.path.join(BASE, f)
    print("### " + f, os.path.getsize(p))
    print(peek(p))
    print()

print("=========== 附件5 结果模板 ===========")
for f in res:
    p = os.path.join(BASE, "附件5", f)
    print("### " + f, os.path.getsize(p))
    print(peek(p))
    print()
